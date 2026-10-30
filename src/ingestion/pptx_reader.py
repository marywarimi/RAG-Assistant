from pathlib import Path
from pptx import Presentation


def load_pptx(pptx_path):
    pptx_path = Path(pptx_path)
    if not pptx_path.exists():
        raise FileNotFoundError(f"Cannot find {pptx_path}")
    presentation = Presentation(str(pptx_path))
    pages = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        text = "\n".join(texts)
        if not text.strip():
            continue
        pages.append({"page": slide_number, "text": text.strip() ,"source": pptx_path.name})
    return pages


if __name__ == "__main__":
    result = load_pptx("Data/papers/raw/LESSON 1 (2).pptx")
    print("Slides with text:", len(result))
    print(result[0]["text"][:100])